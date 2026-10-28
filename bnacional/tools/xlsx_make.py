"""Arma el libro de Excel B_Nacional_historico.xlsx a partir de los CSV generados por bn_export.py.

Uso: python3 -I xlsx_make.py <csv_dir> <salida.xlsx>
Una hoja por CSV (con encabezado fijo y filtros); las tablas ya vienen calculadas.
"""
import csv, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ORDEN = [("tabla_general", "TABLA GENERAL"), ("tabla_temporada", "TABLA POR TEMPORADA"), ("ranking_goleadores", "RANKING GOLEADORES"), ("temporadas", "TEMPORADAS"), ("partidos", "partidos"), ("goleadores", "goleadores"),
         ("reducidos", "reducidos"), ("sanciones", "sanciones"), ("equipos", "equipos"), ("estadios", "estadios"), ("control_fuente", "control fuente"), ("diferencias_entre_fuentes", "diferencias entre fuentes"),
         ("correcciones_aplicadas", "correcciones aplicadas"), ("fuentes", "fuentes"), ("errores_de_lectura", "errores de lectura")]


def main(csv_dir, out):
    wb = Workbook()
    wb.remove(wb.active)
    head = PatternFill("solid", fgColor="17120F")
    for fname, title in ORDEN:
        path = os.path.join(csv_dir, fname + ".csv")
        if not os.path.exists(path):
            continue
        ws = wb.create_sheet(title)
        rows = list(csv.reader(open(path, encoding="utf-8-sig")))
        for r in rows:
            out_r = []
            for v in r:
                if v and v.replace("-", "", 1).replace(".", "", 1).isdigit() and not (len(v) > 1 and v.startswith("0")) and v.count("-") == 0:
                    out_r.append(float(v) if "." in v else int(v))
                else:
                    out_r.append(v)
            ws.append(out_r)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head
            c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i in range(1, ws.max_column + 1):
            width = max((len(str(ws.cell(row=j, column=i).value or "")) for j in range(1, min(ws.max_row, 300) + 1)), default=8)
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
    wb.save(out)
    print(out, os.path.getsize(out) // 1024, "KB")


if __name__ == "__main__":
    main(*sys.argv[1:3])
